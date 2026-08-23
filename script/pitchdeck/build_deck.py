from __future__ import annotations

import json
import math
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches


SCRIPT_ROOT = Path(__file__).resolve().parent
ROOT = SCRIPT_ROOT.parents[1] / "media" / "pitchdeck"
ASSETS = ROOT / "assets"
SLIDES = ROOT / "slides"
BUILD = ROOT / "build"
W, H = 1920, 1080

NAVY = "#06101D"
NAVY_2 = "#0A1B2C"
PANEL = "#0D2236"
PANEL_2 = "#102A41"
WHITE = "#F5FAFF"
MUTED = "#A9BDD0"
CYAN = "#46D9FF"
TEAL = "#35E0C1"
VIOLET = "#9D83FF"
AMBER = "#FFBE62"
RED = "#FF6B73"


def font_path(name: str) -> str:
    candidates = {
        "regular": ["segoeui.ttf", "aptos.ttf", "arial.ttf"],
        "semibold": ["seguisb.ttf", "aptos-semibold.ttf", "arialbd.ttf"],
        "bold": ["segoeuib.ttf", "aptos-bold.ttf", "arialbd.ttf"],
        "mono": ["consola.ttf", "cour.ttf"],
    }[name]
    fonts = Path("C:/Windows/Fonts")
    for candidate in candidates:
        path = fonts / candidate
        if path.exists():
            return str(path)
    return candidates[-1]


FONT_REGULAR = font_path("regular")
FONT_SEMIBOLD = font_path("semibold")
FONT_BOLD = font_path("bold")
FONT_MONO = font_path("mono")


def f(size: int, weight: str = "regular") -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(
        {"regular": FONT_REGULAR, "semibold": FONT_SEMIBOLD, "bold": FONT_BOLD, "mono": FONT_MONO}[weight],
        size,
    )


def hex_rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def gradient_background() -> Image.Image:
    top = hex_rgb(NAVY)
    bottom = hex_rgb("#081828")
    image = Image.new("RGB", (W, H))
    draw = ImageDraw.Draw(image)
    for y in range(H):
        t = y / (H - 1)
        rgb = tuple(round(top[i] * (1 - t) + bottom[i] * t) for i in range(3))
        draw.line((0, y, W, y), fill=rgb)
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    for x in range(0, W, 96):
        od.line((x, 0, x, H), fill=(70, 217, 255, 10), width=1)
    for y in range(0, H, 96):
        od.line((0, y, W, y), fill=(70, 217, 255, 8), width=1)
    return Image.alpha_composite(image.convert("RGBA"), overlay)


def cover_crop(path: Path, box: tuple[int, int, int, int], darken: float = 1.0) -> Image.Image:
    x, y, x2, y2 = box
    target_w, target_h = x2 - x, y2 - y
    img = Image.open(path).convert("RGB")
    ratio = max(target_w / img.width, target_h / img.height)
    size = (round(img.width * ratio), round(img.height * ratio))
    img = img.resize(size, Image.Resampling.LANCZOS)
    left = (img.width - target_w) // 2
    top = (img.height - target_h) // 2
    img = img.crop((left, top, left + target_w, top + target_h))
    if darken != 1.0:
        img = ImageEnhance.Brightness(img).enhance(darken)
    return img.convert("RGBA")


def wrap(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        words = paragraph.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if draw.textlength(candidate, font=font) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def text_block(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    max_width: int,
    spacing: int = 10,
    anchor: str = "la",
) -> int:
    x, y = xy
    lines = wrap(draw, text, font, max_width)
    line_height = font.size + spacing
    for i, line in enumerate(lines):
        draw.text((x, y + i * line_height), line, font=font, fill=fill, anchor=anchor)
    return len(lines) * line_height


def rr(draw: ImageDraw.ImageDraw, box, radius=28, fill=PANEL, outline=None, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def pill(draw: ImageDraw.ImageDraw, xy, text, color=CYAN, width=None):
    x, y = xy
    font = f(25, "semibold")
    w = width or round(draw.textlength(text, font=font) + 48)
    rr(draw, (x, y, x + w, y + 48), radius=24, fill="#0C2940", outline=color, width=2)
    draw.text((x + 24, y + 24), text, font=font, fill=WHITE, anchor="lm")
    return w


def kicker(draw: ImageDraw.ImageDraw, text: str):
    draw.text((92, 64), text.upper(), font=f(24, "semibold"), fill=CYAN)


def title(draw: ImageDraw.ImageDraw, text: str, subtitle: str | None = None, width: int = 1500):
    height = text_block(draw, (92, 112), text, f(58, "bold"), WHITE, width, spacing=10)
    if subtitle:
        text_block(draw, (96, 130 + height), subtitle, f(29), MUTED, width, spacing=9)


def footer(draw: ImageDraw.ImageDraw, number: int, label: str = "CONFIGLAB"):
    draw.text((92, 1030), label, font=f(18, "semibold"), fill="#55738B", anchor="lm")
    draw.text((1828, 1030), f"{number:02d}", font=f(18, "semibold"), fill="#55738B", anchor="rm")


def arrow(draw: ImageDraw.ImageDraw, start, end, color=CYAN, width=5):
    draw.line((*start, *end), fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 18
    for delta in (2.55, -2.55):
        p = (end[0] + length * math.cos(angle + delta), end[1] + length * math.sin(angle + delta))
        draw.line((*end, *p), fill=color, width=width)


def metric(draw, box, value, label, color=CYAN):
    rr(draw, box, fill="#0B2033", outline="#193E59")
    x1, y1, x2, y2 = box
    draw.text((x1 + 28, y1 + 34), value, font=f(55, "bold"), fill=color)
    text_block(draw, (x1 + 30, y1 + 108), label, f(23), MUTED, x2 - x1 - 56, spacing=7)


def slide_1() -> Image.Image:
    img = cover_crop(ASSETS / "configlab-hero.png", (0, 0, W, H), 0.66)
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    od.rectangle((0, 0, 1030, H), fill=(3, 10, 20, 188))
    for x in range(620, 1130):
        alpha = round(188 * (1 - (x - 620) / 510))
        od.line((x, 0, x, H), fill=(3, 10, 20, max(alpha, 0)))
    img = Image.alpha_composite(img, overlay)
    d = ImageDraw.Draw(img)
    pill(d, (92, 84), "RUST • PUBLIC ALPHA", CYAN)
    text_block(d, (92, 205), "ConfigLab", f(92, "bold"), WHITE, 820, spacing=4)
    text_block(d, (98, 330), "Configuration you can explain.\nRuntime change you can trust.", f(45, "semibold"), WHITE, 810, spacing=15)
    text_block(d, (100, 520), "A deterministic, typed configuration foundation for Rust—from startup files to explicit runtime refresh.", f(28), MUTED, 760, spacing=12)
    d.line((100, 735, 590, 735), fill=CYAN, width=4)
    d.text((100, 775), "PITCH DECK  •  AUGUST 2026", font=f(21, "semibold"), fill=CYAN)
    return img


def slide_2() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "The problem"); title(d, "Configuration became critical infrastructure—without a shared model.", "Every project rebuilds the same semantics, then discovers the edge cases in production.")
    sources = [("DEFAULTS", CYAN), ("TOML / JSON / YAML", CYAN), ("ENV", TEAL), ("CLI", VIOLET), ("DEPLOYMENT", AMBER), ("REMOTE", AMBER)]
    y = 420
    for i, (name, color) in enumerate(sources):
        x = 96 + (i % 3) * 270
        yy = y + (i // 3) * 100
        pill(d, (x, yy), name, color, 236)
        arrow(d, (x + 236, yy + 24), (975, 585), color="#365E78", width=3)
    rr(d, (965, 435, 1255, 730), radius=44, fill="#10283D", outline=CYAN, width=3)
    d.text((1110, 510), "?", font=f(105, "bold"), fill=CYAN, anchor="mm")
    d.text((1110, 660), "EFFECTIVE\nCONFIG", font=f(25, "semibold"), fill=WHITE, anchor="mm", align="center", spacing=8)
    consequences = [
        ("UNKNOWN WINNER", "Which source actually won?"),
        ("LATE FAILURE", "Validation happens at use time."),
        ("SECRET LEAKAGE", "Diagnostics expose what they inspect."),
        ("SPLIT LIFECYCLE", "Startup and runtime drift apart."),
    ]
    for i, (head, body) in enumerate(consequences):
        x = 1310
        yy = 390 + i * 140
        rr(d, (x, yy, 1825, yy + 110), fill="#101F31", outline="#243B4E")
        d.ellipse((x + 24, yy + 29, x + 42, yy + 47), fill=RED)
        d.text((x + 60, yy + 26), head, font=f(22, "bold"), fill=WHITE)
        d.text((x + 60, yy + 62), body, font=f(20), fill=MUTED)
    footer(d, 2); return img


def slide_3() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "The insight"); title(d, "Source, applicability, and precedence are three different questions.", "ConfigLab separates them so policy stays explicit and deterministic.")
    cols = [
        ("SOURCE", "Where did it come from?", "file  •  env  •  CLI  •  provider", CYAN),
        ("SELECTOR", "When does it apply?", "environment  •  region  •  tenant", VIOLET),
        ("LAYER", "How much authority?", "base  <  contextual  <  override", AMBER),
    ]
    for i, (head, question, examples, color) in enumerate(cols):
        x = 94 + i * 600
        rr(d, (x, 370, x + 535, 690), fill="#0C2134", outline=color, width=3)
        d.text((x + 36, 410), f"0{i+1}", font=f(22, "bold"), fill=color)
        d.text((x + 36, 468), head, font=f(36, "bold"), fill=WHITE)
        d.text((x + 36, 538), question, font=f(25), fill=MUTED)
        text_block(d, (x + 36, 610), examples, f(22, "mono"), color, 455, spacing=8)
    rr(d, (292, 774, 1628, 926), fill="#0D293D", outline="#2E6078", radius=34)
    d.text((960, 814), "THE RESULT", font=f(20, "bold"), fill=CYAN, anchor="ma")
    d.text((960, 867), "One effective configuration—with a reason for every value.", font=f(34, "semibold"), fill=WHITE, anchor="ma")
    footer(d, 3); return img


def slide_4() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "The product"); title(d, "One semantic pipeline. Every input. One checked result.")
    stages = [
        ("DEFINE", "typed schema\ndefaults\nrequirements", CYAN),
        ("NORMALIZE", "files\nenvironment\narguments", TEAL),
        ("RESOLVE", "layers\nselectors\nmerge + removal", VIOLET),
        ("VALIDATE", "types\nrelationships\ncomponents", AMBER),
        ("CONSUME", "typed value\nprovenance\nrefresh report", CYAN),
    ]
    y1, y2 = 385, 710
    for i, (head, body, color) in enumerate(stages):
        x = 85 + i * 365
        rr(d, (x, y1, x + 290, y2), fill="#0C2032", outline=color, width=3)
        d.ellipse((x + 28, y1 + 30, x + 70, y1 + 72), fill=color)
        d.text((x + 90, y1 + 38), head, font=f(25, "bold"), fill=WHITE)
        text_block(d, (x + 34, y1 + 125), body, f(27), MUTED, 225, spacing=14)
        if i < len(stages) - 1:
            arrow(d, (x + 296, 548), (x + 353, 548), color="#43677F", width=4)
    callouts = ["Source-neutral core", "Derive-first ergonomics", "Runtime is optional"]
    for i, text in enumerate(callouts):
        x = 270 + i * 520
        pill(d, (x, 790), text, [CYAN, VIOLET, AMBER][i], 410)
    footer(d, 4); return img


def slide_5() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "Why it is different"); title(d, "Determinism is a product feature—not an implementation detail.")
    layers = [
        ("override", RED), ("invocation", AMBER), ("environment", VIOLET),
        ("local", TEAL), ("contextual", CYAN), ("base", "#4E87A7"), ("defaults", "#35546B"),
    ]
    base_x, base_y = 125, 875
    for i, (name, color) in enumerate(reversed(layers)):
        yy = base_y - i * 72
        xx = base_x + i * 22
        rr(d, (xx, yy, xx + 650, yy + 54), radius=14, fill="#0E2538", outline=color, width=3)
        d.text((xx + 28, yy + 27), name, font=f(23, "semibold"), fill=WHITE, anchor="lm")
    d.text((180, 330), "EXPLICIT LAYER ORDER", font=f(24, "bold"), fill=CYAN)
    d.text((180, 370), "Authority is declared. Source type never decides the winner.", font=f(22), fill=MUTED)
    guarantees = [
        ("Specificity stays local", "Selectors compare only inside one layer."),
        ("Conflicts fail loudly", "Equal-authority disagreement never depends on input order."),
        ("Schema-owned merges", "Replace, deep merge, append, combine-by-key, and removal are explicit."),
        ("Explain every winner", "Provenance retains defaults, overrides, and shadowed inputs."),
    ]
    for i, (head, body) in enumerate(guarantees):
        x = 955 + (i % 2) * 435
        y = 365 + (i // 2) * 280
        rr(d, (x, y, x + 390, y + 225), fill="#0D2235", outline="#26475D")
        d.text((x + 28, y + 34), f"0{i+1}", font=f(21, "bold"), fill=VIOLET)
        text_block(d, (x + 28, y + 78), head, f(27, "bold"), WHITE, 330, spacing=8)
        text_block(d, (x + 28, y + 136), body, f(21), MUTED, 330, spacing=7)
    footer(d, 5); return img


def slide_6() -> Image.Image:
    img = cover_crop(ASSETS / "runtime-lifecycle.png", (0, 0, W, H), 0.56)
    shade = Image.new("RGBA", (W, H), (0, 0, 0, 0)); sd = ImageDraw.Draw(shade)
    sd.rectangle((0, 0, W, 285), fill=(4, 12, 22, 215)); sd.rectangle((0, 820, W, H), fill=(4, 12, 22, 210))
    img = Image.alpha_composite(img, shade); d = ImageDraw.Draw(img)
    kicker(d, "Runtime without roulette"); title(d, "Acquire. Resolve. Validate. Refresh safely.", "Complete valid values replace the current configuration immediately.")
    labels = [(210, "ACQUIRE", CYAN), (555, "NORMALIZE", CYAN), (930, "RESOLVE", AMBER), (1325, "VALIDATE", VIOLET), (1650, "REFRESH", TEAL)]
    for x, label, color in labels:
        pill(d, (x - 105, 830), label, color, 210)
    d.text((960, 948), "Parse • resolution • validation failure → current value remains unchanged", font=f(26, "semibold"), fill=WHITE, anchor="ma")
    footer(d, 6); return img


def slide_7() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "Trust boundaries"); title(d, "Production-ready means knowing exactly what the library owns.")
    rr(d, (90, 340, 925, 900), fill="#0B2434", outline=TEAL, width=3)
    rr(d, (995, 340, 1830, 900), fill="#211A24", outline=AMBER, width=3)
    d.text((140, 390), "CONFIGLAB OWNS", font=f(27, "bold"), fill=TEAL)
    d.text((1045, 390), "THE HOST OWNS", font=f(27, "bold"), fill=AMBER)
    left = ["Deterministic resolution", "Complete-value validation", "Secret-safe diagnostics", "Transactional source refresh", "Current provenance trace", "External-input reconciliation"]
    right = ["Secret storage and retrieval", "Network delivery and durability", "Distributed rollout coordination", "Rebuilding clients or listeners", "External side-effect atomicity", "Fleet policy and governance"]
    for i, item in enumerate(left):
        y = 480 + i * 62
        d.ellipse((140, y + 4, 158, y + 22), fill=TEAL)
        d.text((182, y), item, font=f(24), fill=WHITE)
    for i, item in enumerate(right):
        y = 480 + i * 62
        d.ellipse((1045, y + 4, 1063, y + 22), outline=AMBER, width=3)
        d.text((1087, y), item, font=f(24), fill=WHITE)
    rr(d, (475, 930, 1445, 992), radius=31, fill="#0E2B3D", outline="#2D5B72")
    d.text((960, 961), "Clear boundaries make integrations safer—and the core reusable.", font=f(25, "semibold"), fill=CYAN, anchor="mm")
    footer(d, 7); return img


def slide_8() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "Developer experience"); title(d, "The simple case stays simple. The hard case has somewhere to grow.")
    rr(d, (90, 340, 1050, 895), fill="#071722", outline="#28516A", width=3)
    d.text((130, 380), "app.rs", font=f(20, "semibold"), fill=MUTED)
    code = [
        ("#[derive(Debug, Config)]", VIOLET),
        ("struct AppConfig {", WHITE),
        ('    #[config(default = 8080, env = "APP_PORT", cli = "port")]', CYAN),
        ("    port: u16,", WHITE),
        ("}", WHITE),
        ("", WHITE),
        ("let config = AppConfig::builder()", WHITE),
        ('    .file(path("config.toml").optional())', TEAL),
        ("    .environment(env())", TEAL),
        ("    .arguments(args())", TEAL),
        ("    .load()?;", AMBER),
    ]
    y = 435
    for line, color in code:
        d.text((135, y), line, font=f(24, "mono"), fill=color)
        y += 39
    features = [
        ("ONE SCHEMA", "TOML, JSON, YAML, environment, CLI"),
        ("TWO MODES", "load once—or compile a reusable runtime plan"),
        ("ZERO PRIVILEGE", "examples consume only the public API"),
        ("OPT-IN COST", "derive and formats are feature-gated"),
    ]
    for i, (head, body) in enumerate(features):
        x = 1120; yy = 345 + i * 135
        rr(d, (x, yy, 1828, yy + 108), fill="#0D2336", outline="#25465D")
        d.text((x + 28, yy + 25), head, font=f(21, "bold"), fill=CYAN)
        d.text((x + 28, yy + 62), body, font=f(22), fill=WHITE)
    d.text((1120, 910), "Ergonomics on top. One resolver underneath.", font=f(28, "semibold"), fill=VIOLET)
    footer(d, 8); return img


def slide_9() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "Proof, not promises"); title(d, "The alpha already behaves like infrastructure.", "Every headline capability is protected by executable contracts and a downstream packaging gate.")
    metric(d, (90, 360, 500, 560), "140", "library + example behavioral tests", CYAN)
    metric(d, (535, 360, 945, 560), "7", "proc-macro unit tests", VIOLET)
    metric(d, (980, 360, 1390, 560), "5", "public-API example applications", TEAL)
    metric(d, (1425, 360, 1830, 560), "1", "semantic resolver", AMBER)
    chain = [
        ("FORMAT", "cargo fmt"), ("TEST", "workspace + docs"), ("LINT", "Clippy -D warnings"),
        ("MATRIX", "isolated features"), ("PACKAGE", "fresh consumer"), ("RELEASE", "optimized probes"),
    ]
    for i, (head, sub) in enumerate(chain):
        x = 90 + i * 292
        rr(d, (x, 675, x + 245, 810), fill="#0B2133", outline="#24475E")
        d.text((x + 22, 708), head, font=f(21, "bold"), fill=CYAN)
        d.text((x + 22, 755), sub, font=f(20), fill=MUTED)
        if i < len(chain) - 1:
            arrow(d, (x + 250, 742), (x + 282, 742), color="#3F647B", width=3)
    d.text((92, 870), "Also enforced: no project unsafe blocks • package MSRV Rust 1.97 • no compiler or rustdoc warnings", font=f(24, "semibold"), fill=WHITE)
    footer(d, 9); return img


def slide_10() -> Image.Image:
    img = gradient_background(); d = ImageDraw.Draw(img)
    kicker(d, "Who wins"); title(d, "A shared configuration contract creates leverage across the organization.")
    personas = [
        ("APPLICATION DEVELOPERS", "Define once. Load typed values. Diagnose quickly.", CYAN),
        ("LIBRARY AUTHORS", "Publish reusable requirements without choosing storage.", VIOLET),
        ("OPERATORS", "See current values, origin, validation failure, and change.", TEAL),
        ("PLATFORM TEAMS", "Standardize semantics before standardizing infrastructure.", AMBER),
    ]
    for i, (head, body, color) in enumerate(personas):
        x = 90 + (i % 2) * 885
        y = 355 + (i // 2) * 285
        rr(d, (x, y, x + 825, y + 235), fill="#0C2235", outline=color, width=3)
        d.ellipse((x + 34, y + 36, x + 86, y + 88), fill=color)
        d.text((x + 115, y + 43), head, font=f(25, "bold"), fill=WHITE)
        text_block(d, (x + 40, y + 125), body, f(26), MUTED, 735, spacing=9)
    rr(d, (420, 915, 1500, 982), radius=34, fill="#0D2A3E", outline="#306179")
    d.text((960, 948), "Start with Rust applications. Grow toward an interoperable configuration ecosystem.", font=f(25, "semibold"), fill=CYAN, anchor="mm")
    footer(d, 10); return img


def slide_11() -> Image.Image:
    img = cover_crop(ASSETS / "ecosystem-roadmap.png", (0, 0, W, H), 0.50)
    shade = Image.new("RGBA", (W, H), (0, 0, 0, 0)); sd = ImageDraw.Draw(shade)
    sd.rectangle((0, 0, W, 285), fill=(4, 12, 22, 220)); sd.rectangle((1050, 280, W, H), fill=(4, 12, 22, 205))
    img = Image.alpha_composite(img, shade); d = ImageDraw.Draw(img)
    kicker(d, "Roadmap"); title(d, "Keep the kernel small. Grow capability around it.")
    phases = [
        ("TODAY", "Typed definitions\nDeterministic resolver\nFiles • env • CLI\nExplainability + live refresh", CYAN),
        ("NEXT", "Schema introspection\nReusable component definitions\nConfiguration testing tools\nProvider adapter patterns", VIOLET),
        ("LATER — VALIDATED DEMAND", "Remote integrations\nDurable revision tooling\nOrganization policy\nControl-plane interoperability", AMBER),
    ]
    for i, (head, body, color) in enumerate(phases):
        x = 1120; y = 330 + i * 225
        rr(d, (x, y, 1830, y + 190), fill="#091B2B", outline=color, width=3)
        d.text((x + 30, y + 28), head, font=f(22, "bold"), fill=color)
        text_block(d, (x + 30, y + 70), body, f(22), WHITE, 635, spacing=7)
    d.text((92, 955), "Roadmap principle: integrations may vary; semantic truth must not.", font=f(27, "semibold"), fill=WHITE)
    footer(d, 11); return img


def slide_12() -> Image.Image:
    img = cover_crop(ASSETS / "configlab-hero.png", (0, 0, W, H), 0.40)
    overlay = Image.new("RGBA", (W, H), (4, 12, 22, 160)); img = Image.alpha_composite(img, overlay)
    d = ImageDraw.Draw(img)
    d.text((960, 115), "CONFIGLAB", font=f(24, "bold"), fill=CYAN, anchor="ma")
    text_block(d, (960, 225), "Turn configuration from glue code\ninto dependable infrastructure.", f(64, "bold"), WHITE, 1400, spacing=16, anchor="ma")
    d.text((960, 450), "The core is working. The next phase is adoption, feedback, and ecosystem proof.", font=f(29), fill=MUTED, anchor="ma")
    asks = [
        ("ADOPT", "Run ConfigLab in real Rust services."),
        ("CHALLENGE", "Bring the configuration edge cases that break ordinary loaders."),
        ("BUILD", "Shape reusable components and provider integrations."),
    ]
    for i, (head, body) in enumerate(asks):
        x = 150 + i * 570
        rr(d, (x, 590, x + 500, 790), fill="#0B2033", outline=[CYAN, VIOLET, AMBER][i], width=3)
        d.text((x + 30, 630), f"0{i+1}  {head}", font=f(25, "bold"), fill=[CYAN, VIOLET, AMBER][i])
        text_block(d, (x + 30, 690), body, f(24), WHITE, 430, spacing=8)
    pill(d, (645, 875), "CONFIGURATION YOU CAN EXPLAIN", CYAN, 630)
    d.text((960, 992), "BSD-3-CLAUSE  •  CONFIGLAB 0.1.0-ALPHA.1", font=f(20, "semibold"), fill=MUTED, anchor="ma")
    return img


SLIDE_BUILDERS = [slide_1, slide_2, slide_3, slide_4, slide_5, slide_6, slide_7, slide_8, slide_9, slide_10, slide_11, slide_12]

TITLES = [
    "ConfigLab",
    "Configuration is fragmented",
    "Separate source, selector, and layer",
    "One semantic pipeline",
    "Determinism is a feature",
    "Safe runtime refresh",
    "Explicit trust boundaries",
    "Simple developer experience",
    "Proof, not promises",
    "Who wins",
    "Small core, growing ecosystem",
    "The ask",
]

NARRATION = [
    "Configuration sits on every critical path, yet most teams still treat it as glue code. ConfigLab makes configuration deterministic, typed, explainable, and safe to evolve—from ordinary startup files to deliberate runtime change.",
    "Applications combine defaults, files, environment variables, command-line flags, deployment systems, and eventually remote providers. Without one model, teams cannot confidently answer which value won, why it won, whether it is valid, or whether it can change safely while the process is running.",
    "ConfigLab starts with a clean separation. Source answers where input came from. Selector answers when it applies. Layer answers how much authority it has. Once those concerns stop leaking into one another, resolution becomes predictable and every effective value can carry an explanation.",
    "The product is one semantic pipeline. Applications define a typed contract. Physical sources normalize into source-neutral logical inputs. One resolver applies layers, selectors, merge and removal rules. Complete-value validation runs before the application receives either a typed value or a refresh report.",
    "Determinism is the foundation. Source type never secretly decides precedence. Selector specificity only compares peers inside one layer. Equal-authority conflicts fail loudly. Collection behavior belongs to the schema. Provenance explains defaults, winners, overrides, and shadowed inputs.",
    "Runtime configuration uses the same plan and resolver as startup. A refresh rereads every registered physical source, combines any retained source-neutral external inputs, and validates the complete result. A successful refresh immediately replaces the current typed value; a parse, resolution, or validation failure leaves it unchanged.",
    "Production trust comes from boundaries. ConfigLab owns resolution, validation, redaction, transactional in-process refresh, current provenance, and redaction-aware change reports. The host decides when to refresh and owns remote transport, secret retrieval, distributed rollout, history, and arbitrary external side effects.",
    "The common path remains small: derive a typed configuration, add an optional file, environment, and arguments, then load. The same schema handles TOML, JSON, and YAML. Long-running applications retain those descriptors and refresh the current value without introducing a second semantic system.",
    "This is not a slide-only architecture. The repository carries one hundred forty library and example behavioral tests, seven proc-macro unit tests, five public-API applications, and exactly one semantic resolver. The gate denies warnings, checks isolated features, builds release artifacts, and compiles a fresh packaged consumer.",
    "The leverage compounds. Application developers remove bespoke parsing and validation. Library authors publish source-neutral requirements. Operators get effective state and provenance. Platform teams standardize configuration semantics before committing to any one storage or control-plane technology.",
    "The roadmap protects the core. Today delivers typed definitions, deterministic resolution, explainability, and explicit runtime refresh. Next comes introspection, reusable components, testing tools, and provider patterns. Remote transports and control planes remain separate projects when real adoption validates the demand.",
    "The invitation is simple: adopt the alpha in real Rust services, challenge the model with the edge cases that break ordinary loaders, and help shape reusable components and integrations. ConfigLab can turn configuration from repeated glue code into dependable infrastructure.",
]


def build_slides() -> list[Path]:
    SLIDES.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for index, builder in enumerate(SLIDE_BUILDERS, 1):
        image = builder().convert("RGB")
        path = SLIDES / f"{index:02d}-{TITLES[index-1].lower().replace(' ', '-').replace(',', '')}.png"
        image.save(path, "PNG", optimize=True)
        paths.append(path)
    return paths


def build_pptx(paths: list[Path]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path, title_text, notes in zip(paths, TITLES, NARRATION):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        try:
            notes_frame = slide.notes_slide.notes_text_frame
            notes_frame.text = f"{title_text}\n\n{notes}"
        except AttributeError:
            pass
    prs.core_properties.title = "ConfigLab — Configuration You Can Explain"
    prs.core_properties.subject = "Product and engineering pitch deck"
    prs.core_properties.author = "Dominic Maabobra Tuolong"
    prs.core_properties.comments = "Generated from repository evidence; speaker narration is included in slide notes when supported."
    out = ROOT / "ConfigLab-Pitch-Deck.pptx"
    prs.save(out)
    return out


def build_pdf(paths: list[Path]) -> Path:
    pages = [Image.open(path).convert("RGB") for path in paths]
    out = ROOT / "ConfigLab-Pitch-Deck.pdf"
    pages[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=pages[1:], quality=92)
    for page in pages:
        page.close()
    return out


def build_manifest(paths: list[Path]) -> Path:
    manifest = {
        "title": "ConfigLab — Configuration You Can Explain",
        "audience": "Engineering leaders, platform teams, Rust adopters, and ecosystem sponsors",
        "slides": [
            {"number": i, "title": title, "image": str(path.relative_to(ROOT)), "narration": narration}
            for i, (title, path, narration) in enumerate(zip(TITLES, paths, NARRATION), 1)
        ],
    }
    out = BUILD / "deck-manifest.json"
    out.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    script = ROOT / "SCRIPT.md"
    chunks = ["# ConfigLab pitch video script\n"]
    for i, (title, narration) in enumerate(zip(TITLES, NARRATION), 1):
        chunks.append(f"## {i:02d}. {title}\n\n{narration}\n")
    script.write_text("\n".join(chunks), encoding="utf-8")
    return out


def build_video(paths: list[Path], manifest: Path) -> Path | None:
    narration_dir = BUILD / "narration"
    narration_dir.mkdir(parents=True, exist_ok=True)
    synth = SCRIPT_ROOT / "build_narration.ps1"
    subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(synth), "-Manifest", str(manifest), "-Output", str(narration_dir)],
        check=True,
    )
    try:
        import imageio_ffmpeg
    except ImportError:
        print("imageio-ffmpeg is not installed; skipping video", file=sys.stderr)
        return None
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    segments = BUILD / "segments"
    if segments.exists():
        shutil.rmtree(segments)
    segments.mkdir(parents=True)
    segment_paths = []
    for index, image_path in enumerate(paths, 1):
        audio = narration_dir / f"{index:02d}.wav"
        segment = segments / f"{index:02d}.mp4"
        subprocess.run(
            [
                ffmpeg, "-y", "-loop", "1", "-framerate", "30", "-i", str(image_path), "-i", str(audio),
                "-vf", "fade=t=in:st=0:d=0.35,fade=t=out:st=100000:d=0.35",
                "-c:v", "libx264", "-preset", "medium", "-crf", "19", "-pix_fmt", "yuv420p",
                "-c:a", "aac", "-b:a", "160k", "-shortest", "-movflags", "+faststart", str(segment),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        segment_paths.append(segment)
    concat = BUILD / "segments.txt"
    concat.write_text("\n".join(f"file '{p.as_posix()}'" for p in segment_paths), encoding="utf-8")
    out = ROOT / "ConfigLab-Pitch-Video.mp4"
    subprocess.run(
        [ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(concat), "-c", "copy", "-movflags", "+faststart", str(out)],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    return out


def main() -> None:
    BUILD.mkdir(parents=True, exist_ok=True)
    paths = build_slides()
    pptx = build_pptx(paths)
    pdf = build_pdf(paths)
    manifest = build_manifest(paths)
    video = build_video(paths, manifest)
    print(f"Built {len(paths)} slides")
    print(pptx)
    print(pdf)
    if video:
        print(video)


if __name__ == "__main__":
    main()
