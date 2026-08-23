from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import imageio_ffmpeg
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2] / "media" / "pitchdeck"
slides = sorted((ROOT / "slides").glob("*.png"))
assert len(slides) == 12, f"expected 12 slide images, found {len(slides)}"

for slide in slides:
    with Image.open(slide) as image:
        assert image.size == (1920, 1080), f"unexpected size for {slide.name}: {image.size}"

prs = Presentation(ROOT / "ConfigLab-Pitch-Deck.pptx")
assert len(prs.slides) == 12, f"expected 12 PowerPoint slides, found {len(prs.slides)}"
notes = []
for slide in prs.slides:
    try:
        notes.append(slide.notes_slide.notes_text_frame.text.strip())
    except AttributeError:
        notes.append("")
assert all(notes), "every PowerPoint slide must include speaker notes"

pdf = PdfReader(ROOT / "ConfigLab-Pitch-Deck.pdf")
assert len(pdf.pages) == 12, f"expected 12 PDF pages, found {len(pdf.pages)}"

ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
probe = subprocess.run(
    [ffmpeg, "-i", str(ROOT / "ConfigLab-Pitch-Video.mp4")],
    text=True,
    stdout=subprocess.PIPE,
    stderr=subprocess.PIPE,
)
match = re.search(r"Duration: (\d+):(\d+):(\d+\.\d+)", probe.stderr)
assert match, "could not determine video duration"
hours, minutes, seconds = match.groups()
duration = int(hours) * 3600 + int(minutes) * 60 + float(seconds)
assert duration > 120, f"video is unexpectedly short: {duration} seconds"

thumb_w, thumb_h = 480, 270
sheet = Image.new("RGB", (thumb_w * 4, thumb_h * 3), "#06101D")
for index, slide in enumerate(slides):
    with Image.open(slide).convert("RGB") as image:
        image.thumbnail((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        x = (index % 4) * thumb_w
        y = (index // 4) * thumb_h
        sheet.paste(image, (x, y))
sheet.save(ROOT / "ConfigLab-Pitch-Deck-Contact-Sheet.png", "PNG")

report = {
    "slide_images": len(slides),
    "slide_resolution": "1920x1080",
    "powerpoint_slides": len(prs.slides),
    "powerpoint_slides_with_notes": sum(bool(note) for note in notes),
    "pdf_pages": len(pdf.pages),
    "video_duration_seconds": round(duration, 2),
    "video_codec_probe": "h264" if "Video: h264" in probe.stderr else "unknown",
    "video_audio_probe": "aac" if "Audio: aac" in probe.stderr else "unknown",
}
(ROOT / "validation-report.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
print(json.dumps(report, indent=2))
